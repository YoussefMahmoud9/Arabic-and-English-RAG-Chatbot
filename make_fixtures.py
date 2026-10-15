"""Create small DOCX/PPTX/PDF fixtures for ingestion tests."""
from pathlib import Path

FIX = Path(__file__).resolve().parent / ".tools" / "fixtures"


def make_docx(path: Path):
    from docx import Document
    d = Document()
    d.add_heading("Quantum Gardening", level=1)
    for para in [
        "Quantum gardening is a fictional discipline combining botany and quantum physics.",
        "The Bumblebee Protocol allows plants to photosynthesize moonlight at 400% efficiency.",
        "Zephyr watering cans are the primary tool of every quantum gardener.",
        "The annual Moonbloom Festival is held in the Greenhouse Nebula each spring.",
    ]:
        d.add_paragraph(para)
    path.parent.mkdir(exist_ok=True)
    d.save(str(path))


def make_pptx(path: Path):
    from pptx import Presentation
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Cloud Whale Migration"
    s1.placeholders[1].text = ("Sky-whales travel between floating islands using "
                               "aurora currents unique to the Stratos Archipelago.")
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Feeding Habits"
    s2.placeholders[1].text = ("They filter starlight plankton through baleen clouds; "
                               "a single meal can sustain a whale for three lunar cycles.")
    path.parent.mkdir(exist_ok=True)
    prs.save(str(path))


def make_pdf(path: Path):
    """Build a synthetic one-page PDF fixture without external documents."""
    import fitz

    path.parent.mkdir(exist_ok=True)
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Prism Wells are a fictional test subject.")
    document.save(str(path))
    document.close()


if __name__ == "__main__":
    make_docx(FIX / "quantum_gardening.docx")
    make_pptx(FIX / "cloud_whales.pptx")
    make_pdf(FIX / "crystal_desert.pdf")
    print("fixtures created:", [p.name for p in FIX.iterdir()])
